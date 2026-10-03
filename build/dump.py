import sys
from pptx import Presentation
p = Presentation(sys.argv[1]); S = list(p.slides)
for i in map(int, sys.argv[2].split(',')):
    print(f"=== slide {i}")
    def walk(shapes, ind=''):
        for sh in shapes:
            if sh.shape_type == 6: walk(sh.shapes, ind+'  '); continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                print(f"{ind}[{sh.shape_id}] " + " ¶ ".join(pp.text for pp in sh.text_frame.paragraphs))
            elif getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows: print(f"{ind}[{sh.shape_id}] ROW: " + " | ".join(c.text for c in r.cells))
            elif sh.shape_type == 13: print(f"{ind}[{sh.shape_id}] PICTURE {sh.name}")
    walk(s:=S[i].shapes)
