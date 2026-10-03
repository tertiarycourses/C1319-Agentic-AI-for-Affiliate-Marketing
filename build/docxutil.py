import re
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
OLD_C = "TGS-2025060552"; NEW_C = "C1319"
OLD_URL = "https://www.tertiarycourses.com.sg/wsq-agentic-ai-for-affiliate-marketing.html"
NEW_URL = "https://www.tertiarycourses.com.sg/agentic-ai-for-affiliate-marketing.html"
DROPPED = [3, 10, 11, 13, 14, 15, 16, 195, 197, 198]   # 1-based source slides removed from the deck
LAB_MIN = {1: 25, 2: 55, 3: 35, 4: 35, 5: 55, 6: 55}

def remap(n):
    assert n not in DROPPED, n
    return n - sum(1 for d in DROPPED if d < n)

def set_para(pp, text):
    runs = [r for r in pp.runs if r.text] or pp.runs
    if not runs: pp.add_run(text); return
    runs[0].text = text
    for r in runs[1:]: r.text = ''

def set_cell(cell, text):
    set_para(cell.paragraphs[0], text)
    for pp in cell.paragraphs[1:]: pp._p.getparent().remove(pp._p)

def remove(el): el.getparent().remove(el)

def all_paras(doc):
    for pp in doc.paragraphs: yield pp
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells: yield from c.paragraphs
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer, s.even_page_footer):
            yield from part.paragraphs

def global_subs(doc, extra=()):
    subs = list(extra) + [(OLD_URL, NEW_URL), ("TGS Ref No: " + OLD_C, "Course Code: " + NEW_C),
                          (OLD_C, NEW_C), ("Version 2.0", "Version 1.0")]
    for pp in all_paras(doc):
        new = pp.text
        for a, b in subs: new = new.replace(a, b)
        if new != pp.text: set_para(pp, new)

def version_table(t, summary):
    for r in t.rows[1:-1]: remove(r._tr)
    for c, v in zip(t.rows[-1].cells, ["1.0", "3 October 2026", summary, "Dr Alfred Ang"]): set_cell(c, v)

def drop_section(doc, heading, style, keep_tables=False):
    """Remove a heading and everything up to the next heading of the same or higher level."""
    lvl = int(style[-1])
    body = list(doc.element.body.iterchildren())
    start = next(i for i, el in enumerate(body) if el.tag.endswith('}p') and Paragraph(el, None).text == heading)
    out = [body[start]]
    for el in body[start + 1:]:
        if keep_tables and el.tag.endswith('}tbl'): break
        if el.tag.endswith('}p'):
            ps = el.find('.//' + qn('w:pStyle'))
            sv = ps.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') if ps is not None else ''
            m = re.fullmatch(r'Heading(\d)', sv or '')
            if m and int(m.group(1)) <= lvl: break
            if el.find('.//' + qn('w:sectPr')) is not None: break
        out.append(el)
    for el in out: remove(el)
    return len(out)

def drop_toc(doc, title):
    for pp in doc.paragraphs:
        if pp.text.split('\t')[0] == title and '\t' in pp.text: remove(pp._p); return
    raise KeyError(title)

def fix_toc(doc, pages):
    for pp in doc.paragraphs:
        if '\t' in pp.text:
            title = pp.text.split('\t')[0]
            if title in pages: set_para(pp, f"{title}\t{pages[title]}")
