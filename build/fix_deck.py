"""C1319 deck: post-convert_deck fixes (cover, admin cards, one-day schedule, LO eyebrows, lab minutes, close, Course Feedback)."""
import re
from pptx import Presentation

F = "courseware/Agentic AI for Affiliate Marketing-v1.0.pptx"
p = Presentation(F); S = list(p.slides)
LAB_MIN = {1: 25, 2: 55, 3: 35, 4: 35, 5: 55, 6: 55}   # LP one-day schedule (x0.6 of the WSQ teaching time)

def shape(slide, sid): return next(s for s in slide.shapes if s.shape_id == sid)
def set_text(sh, text, expect=None):
    if expect is not None: assert sh.text_frame.text == expect, (sh.text_frame.text, expect)
    pars = sh.text_frame.paragraphs
    runs = [r for r in pars[0].runs if r.text] or pars[0].runs
    runs[0].text = text
    for r in runs[1:]: r.text = ''
    for pp in pars[1:]:
        for r in pp.runs: r.text = ''
def put(slide, sid, text, expect): set_text(shape(slide, sid), text, expect)

put(S[0], 6, 'NON-WSQ COURSEWARE', 'WSQ COURSEWARE')
# Ground rules
put(S[5], 38, 'Work hands-on', 'Work independently')
put(S[5], 39, 'Build and verify your own lab evidence', 'Assessment evidence must be your own')
put(S[5], 44, 'Full attendance', 'Attendance')
put(S[5], 45, 'Attend the whole day to get the most from the course', 'At least 75% attendance is required')
put(S[6], 28, 'Keep materials for reference after the course', 'Keep materials available for the open-book assessment')
# Schedule -> one day
put(S[7], 4, 'SCHEDULE · 7.5 HOURS', 'SCHEDULE · 16 HOURS')
put(S[7], 5, 'One-Day Learning Journey', 'Two-Day Learning Journey')
put(S[7], 14, 'Morning', 'Day 1 AM')
put(S[7], 20, 'Late morning', 'Day 1 PM')
put(S[7], 26, 'Afternoon', 'Day 2 AM')
put(S[7], 32, 'Late afternoon', 'Day 2 PM')
put(S[7], 33, 'Commission reconciliation, reporting and optimisation',
    'Commission reconciliation, reporting, optimisation and assessment')
put(S[13], 33, 'audit artifact', 'audit and assessment artifact')
# Course close
put(S[184], 5, 'From Evidence to Practice', 'From Evidence to Assessment')
put(S[184], 6, 'Review the control loop and preserve your working artifacts',
    'Review the control loop, complete the assessment and preserve your working artifacts')
# Final attendance -> Course Feedback (rewritten in place; sits between Support and Thank You)
fb = S[188]
put(fb, 4, 'COURSE FEEDBACK', 'FINAL ATTENDANCE')
put(fb, 5, 'Course Feedback', 'Digital Attendance (Mandatory)')
put(fb, 14, 'Trainer', 'AM')
put(fb, 15, 'Rate the delivery, pace and clarity of the explanations', 'Attendance recorded')
put(fb, 20, 'Materials', 'PM')
put(fb, 21, 'Tell us how the slides, Learner Guide and labs worked for you', 'Attendance recorded')
put(fb, 26, 'Improve', 'Assessment')
put(fb, 27, 'Suggest one change that would make the next class better', 'Complete the final SSG attendance event')
# LO eyebrows: keep LO#, strip K#/A# criteria; lab eyebrows get the one-day minutes
n_lo = n_lab = 0
for i, s in enumerate(S):
    title = ' '.join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if re.fullmatch(r'LO\d \| [KA0-9, |]+', t):
            set_text(sh, t[:3]); n_lo += 1
        m = re.fullmatch(r'(LO\d) \| [KA0-9, |]+ · \d+ minutes', t)
        if m:
            lab = int(re.search(r'\bLAB (\d)\b', title).group(1))
            set_text(sh, f'{m.group(1)} · {LAB_MIN[lab]} minutes'); n_lab += 1
assert (n_lo, n_lab) == (3, 6), (n_lo, n_lab)
p.save(F)
print('ok', n_lo, n_lab, len(S))
