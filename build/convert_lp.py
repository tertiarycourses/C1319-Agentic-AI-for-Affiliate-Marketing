"""C1319 Lesson Plan from the WSQ LP: strip assessment/attendance/TSC, collapse to ONE day 9:30am-5:30pm
(450 min teaching + one 30-min lunch). All 3 topics and 6 labs kept in the parent's order, rescaled x0.6."""
import copy, re, sys
from docx import Document
from docx.table import _Row
from docx.text.paragraph import Paragraph
sys.path.insert(0, 'build'); from docxutil import *
SRC = "source-built/LP-Agentic AI for Affiliate Marketing.docx"
OUT = "courseware/LP-Agentic AI for Affiliate Marketing.docx"
d = Document(SRC)
DAY1_OLD = "Day 1 - Research markets, evaluate offers and build the partner portfolio"
DAY1_NEW = "Day 1 · 9:30am – 5:30pm - Research offers, build the partner portfolio, run controlled campaigns and report"
global_subs(d, [(DAY1_OLD, DAY1_NEW), ("Two-Day Schedule", "One-Day Schedule"),
                ("Lab and Assessment Traceability", "Lab Traceability"),
                ("Use lab acceptance checklists as formative assessment and remediation guides.",
                 "Use lab acceptance checklists as progress checks and remediation guides.")])
T = list(d.tables)   # captured before any table is removed
version_table(T[0], "One-day edition (9:30am – 5:30pm, 7.5 instructional hours) aligned to six connected labs.")

drop_toc(d, "Assessment Strategy")
drop_toc(d, "Day 2 - Create controlled campaigns, reconcile results and report improvements")
assert drop_section(d, "Assessment Strategy", "Heading1") == 4

# Course information
info = T[1]
for r in list(info.rows):
    k = r.cells[0].text
    if k == "TSC": remove(r._tr)
    if k == "Duration": set_cell(r.cells[1], "1 day - 7.5 instructional hours")
    if k == "Timing": set_cell(r.cells[1], "9:30am – 5:30pm; 30-minute lunch break; refreshments taken during activity time")

# One-day schedule: Day 1 table receives the Day 2 teaching rows; Day 2 heading/table/total removed.
t1, t2 = T[2], T[3]
r1, r2 = list(t1.rows), list(t2.rows)
plan = [  # (row element, activity text or None) in delivery order
    (r1[1]._tr, "Welcome, outcomes and learning contract"),
    (r1[2]._tr, None), (r1[4]._tr, None), (r1[5]._tr, None), (r1[7]._tr, None),
    ("LUNCH", None),
    (r1[9]._tr, None), (r1[10]._tr, None),
    (copy.deepcopy(r2[2]._tr), None), (copy.deepcopy(r2[4]._tr), None),
    (copy.deepcopy(r2[5]._tr), None), (copy.deepcopy(r2[7]._tr), None),
    (copy.deepcopy(r2[9]._tr), "Course synthesis, Q&A and course feedback"),
]
lunch = r1[6]._tr
for r in r1[1:]: remove(r._tr)
teach = [tr for tr, _ in plan if tr != "LUNCH"]
src = [int(_Row(tr, t1).cells[1].text) for tr in teach]
mins = [round(m * 450 / sum(src) / 5) * 5 for m in src]
mins[mins.index(max(mins))] += 450 - sum(mins)
assert sum(mins) == 450, mins
def hm(m): return f"{m // 60}:{m % 60:02d}"
t = 9 * 60 + 30; k = 0; tbl = t1._tbl
for tr, act in plan:
    if tr == "LUNCH":
        tbl.append(lunch); row = _Row(lunch, t1)
        set_cell(row.cells[0], f"{hm(t)}-{hm(t + 30)}"); set_cell(row.cells[1], "30"); t += 30; continue
    tbl.append(tr); row = _Row(tr, t1)
    set_cell(row.cells[0], f"{hm(t)}-{hm(t + mins[k])}"); set_cell(row.cells[1], str(mins[k]))
    if act: set_cell(row.cells[3], act)
    t += mins[k]; k += 1
assert t == 17 * 60 + 30, hm(t)
print("src", src, "->", mins)
# Day 2 heading + its table + its training-time line
day2 = next(pp for pp in d.paragraphs if pp.text.startswith("Day 2 - "))
nxt = day2._p.getnext(); nxt2 = nxt.getnext()
assert nxt is t2._tbl and Paragraph(nxt2, None).text.startswith("Training time")
for el in (day2._p, nxt, nxt2): remove(el)
for pp in d.paragraphs:
    if pp.text.startswith("Training time: 480 minutes"):
        set_para(pp, "Day total: 7.5 instructional hours (9:30am – 5:30pm), plus a 30-minute lunch.")

# Lab traceability: LO only, drop the assessment-bridge column
tr_t = T[4]
set_cell(tr_t.rows[0].cells[1], "Outcome")
for r in tr_t.rows[1:]: set_cell(r.cells[1], r.cells[1].text.split(' |')[0])
from docx.oxml.ns import qn
grid = tr_t._tbl.tblGrid; remove(grid.findall(qn('w:gridCol'))[3])
for r in tr_t.rows: remove(r._tr.findall(qn('w:tc'))[3])

# Slide reference map: remap through the kept-slide map
sm = T[5]
for r in sm.rows[1:]:
    a, b = map(int, r.cells[1].text.split('-'))
    set_cell(r.cells[1], f"{remap(a)}-{remap(b)}")
    lbl = r.cells[0].text
    if lbl == "Administration and assessment briefing": set_cell(r.cells[0], "Administration")
    if lbl == "Revision, assessment and close": set_cell(r.cells[0], "Revision and close")
    if "From Evidence to Assessment" in r.cells[2].text:
        set_cell(r.cells[2], r.cells[2].text.replace("From Evidence to Assessment", "From Evidence to Practice"))
d.core_properties.title = "Agentic AI for Affiliate Marketing"
d.save(OUT); print("wrote", OUT)
